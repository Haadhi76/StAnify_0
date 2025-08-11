from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse

import requests
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT
from pptx.util import Cm, Pt

from ..chunks_contracts import ChunkSpec
from ..manifest_contracts import RunManifest


def _overall_score(scores):
    """
    Calculate overall score from CritiqueScores (works with dict or Pydantic model).
    
    Args:
        scores: CritiqueScores model or dict
        
    Returns:
        float: Overall score or None if cannot calculate
    """
    if not scores:
        return None
    
    # if it's a dict already
    if isinstance(scores, dict):
        if "overall_score" in scores:
            return scores["overall_score"]
        keys = ("alignment", "vocab", "scope", "cognitive_load")
        vals = [scores.get(k) for k in keys if k in scores and isinstance(scores.get(k), (int, float))]
        return round(sum(vals) / len(vals), 2) if vals else None
    
    # otherwise assume Pydantic model (v1)
    try:
        a = getattr(scores, "alignment", None)
        v = getattr(scores, "vocab", None)
        s = getattr(scores, "scope", None)
        c = getattr(scores, "cognitive_load", None)
        vals = [x for x in (a, v, s, c) if isinstance(x, (int, float))]
        return round(sum(vals) / len(vals), 2) if vals else None
    except Exception:
        return None


def _fetch_image(uri: str) -> BytesIO:
    if not uri:
        raise FileNotFoundError("Empty image URI")

    # file:// support (Windows & *nix)
    if uri.startswith("file://"):
        parsed = urlparse(uri)
        p = Path(unquote(parsed.path))
        if parsed.netloc:  # windows drive in netloc (rare), join it
            p = Path(f"{parsed.netloc}{p.as_posix()}")
        with open(p, "rb") as f:
            data = f.read()
        bio = BytesIO(data); bio.seek(0); return bio

    # http(s)
    if uri.startswith("http://") or uri.startswith("https://"):
        r = requests.get(uri, timeout=30)
        r.raise_for_status()
        bio = BytesIO(r.content); bio.seek(0); return bio

    # plain local path
    with open(Path(uri), "rb") as f:
        data = f.read()
    bio = BytesIO(data); bio.seek(0); return bio

def export_manifest_to_pptx(manifest: RunManifest, out_path: str) -> str:
    prs = Presentation()
    
    # === COVER SLIDE ===
    cover_slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title slide layout
    cover_slide.shapes.title.text = f"🎓 {manifest.task.subject.title()} — {manifest.task.year}"
    
    # Cover subtitle with topic and objectives
    subtitle_text = f"Topic: {manifest.refined_prompt.topic_title}\n\n"
    subtitle_text += "📋 Learning Objectives:\n"
    
    # Extract or create objectives
    objectives = getattr(manifest.refined_prompt, 'objectives', None)
    if not objectives:
        objectives = [
            f"Understand {manifest.refined_prompt.topic_title.lower()}",
            "Apply learning through visual examples",
            "Demonstrate understanding through practice"
        ]
    
    for obj in objectives:
        subtitle_text += f"• {obj}\n"
    
    subtitle_text += f"\nRun ID: {manifest.task.run_id}"
    cover_slide.placeholders[1].text = subtitle_text
    
    # === CONTENT SLIDES ===
    chunk_by_id = {c.chunk_id: c for c in [ChunkSpec(**cc) if isinstance(cc, dict) else cc for cc in manifest.chunks]}

    for panel in sorted(manifest.panels, key=lambda p: p.index):
        s = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Title bar
        title_text = f"Panel {panel.index + 1}: {chunk_by_id.get(panel.chunk_id).title if chunk_by_id.get(panel.chunk_id) else panel.chunk_id}"
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), Cm(0.5), prs.slide_width - Cm(1), Cm(1.2))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(240,240,240)
        tx = bar.text_frame; tx.clear(); p = tx.paragraphs[0]; p.text = title_text; p.font.size = Pt(18)

        # Image
        left = Cm(1.0); top = Cm(2.2); right_margin = Cm(1.0); cap_h = Cm(2.0)
        area_w = prs.slide_width - left - right_margin
        area_h = prs.slide_height - top - cap_h - Cm(0.8)
        try:
            bio = _fetch_image(panel.image_uri)
            pic = s.shapes.add_picture(bio, left, top)
            scale = min(area_w / pic.width, area_h / pic.height)
            pic.width = int(pic.width * scale); pic.height = int(pic.height * scale)
            pic.left = int(left + (area_w - pic.width) / 2)
        except Exception:
            ph = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, area_w, area_h)
            ph.text = "Image unavailable"

        # Caption
        cap = s.shapes.add_textbox(Cm(1.0), prs.slide_height - cap_h - Cm(0.5), prs.slide_width - Cm(2.0), cap_h)
        tf = cap.text_frame; tf.clear(); q = tf.paragraphs[0]; q.text = panel.caption or ""; q.font.size = Pt(18); q.alignment = PP_PARAGRAPH_ALIGNMENT.CENTER

    # === CLOSING SLIDE ===
    closing_slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and content layout
    closing_slide.shapes.title.text = "✅ Success Criteria & Assessment"
    
    # Success criteria content
    success_text = "🎯 Success Criteria:\n"
    success_criteria = [
        f"Can explain key concepts from {manifest.refined_prompt.topic_title.lower()}",
        "Demonstrates understanding through examples",
        "Shows confidence in applying new knowledge",
        "Can connect learning to real-world situations"
    ]
    
    for criteria in success_criteria:
        success_text += f"✓ {criteria}\n"
    
    # Add critique information if available
    if manifest.overall_critique:
        critique = manifest.overall_critique
        verdict_emoji = "🟢" if critique.verdict == "Accept" else "🟡" if critique.verdict == "Revise" else "🔴"
        
        success_text += f"\n📊 Content Quality Assessment:\n"
        success_text += f"Overall Verdict: {verdict_emoji} {critique.verdict}\n"
        
        if hasattr(critique, 'scores') and critique.scores:
            scores = critique.scores
            success_text += f"Quality Scores: "
            success_text += f"Clarity: {getattr(scores, 'clarity', 'N/A')} | "
            success_text += f"Engagement: {getattr(scores, 'engagement', 'N/A')} | "
            success_text += f"Accuracy: {getattr(scores, 'accuracy', 'N/A')}\n"
        
        if critique.evidence:
            success_text += f"\nKey Feedback:\n"
            for evidence in critique.evidence[:2]:  # Top 2 feedback points for space
                success_text += f"• {evidence}\n"
    
    success_text += f"\n🤖 Generated by StAnify AI\nEducational content platform with quality assurance"
    
    closing_slide.placeholders[1].text = success_text

    prs.save(out_path)
    return out_path
